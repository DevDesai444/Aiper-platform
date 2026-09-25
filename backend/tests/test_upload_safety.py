"""The upload path against hostile input.

Every scenario here is an attack that used to work (or would have): a
traversal name, a spoofed extension, a body that lies about its size, a file
that parses forever. The properties asserted are the ones that matter — the
API answers, nothing user-controlled reaches a path, no failure leaves bytes
or rows behind — not the incidental wording of messages.

Real Postgres, real router, real parse sandbox (spawned children); the
vector store is a recording fake, because indexing needs Azure and these
tests are about what happens before and around it.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
import pytest_asyncio
from app.config import settings
from app.core.body_limit import BodySizeLimitMiddleware
from app.db.models import FileAsset
from app.rag.loaders import ParseRejected
from app.rag.sandbox import ParseTimeout, parse_in_sandbox
from app.services.uploads import sanitize_filename
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import func, select

from .factories import make_grant, make_org, make_project, make_user

# ────────────────────────────── file builders ──────────────────────────────


def minimal_pdf(text: str = "Hello from a real PDF page.") -> bytes:
    """A handcrafted one-page PDF with extractable text and a valid xref."""

    def obj(num: int, body: str) -> str:
        return f"{num} 0 obj\n{body}\nendobj\n"

    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        obj(1, "<< /Type /Catalog /Pages 2 0 R >>"),
        obj(2, "<< /Type /Pages /Kids [3 0 R] /Count 1 >>"),
        obj(
            3,
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        ),
        f"4 0 obj\n<< /Length {len(stream)} >>\nstream\n{stream}\nendstream\nendobj\n",
        obj(5, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for piece in objects:
        offsets.append(len(out))
        out += piece
    xref_pos = len(out)
    out += "xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n"
    out += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
    return out.encode("latin-1")


def docx_bytes(paragraphs: tuple[str, ...] = ("Systems shall be verified.",)) -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: int = 2) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    buf = io.BytesIO()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for index in range(slides):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Slide {index + 1} content"
    prs.save(buf)
    return buf.getvalue()


def plain_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("readme.txt", "just a zip, not an office file")
    return buf.getvalue()


# ─────────────────────────────── app fixture ───────────────────────────────


class FakeStore:
    """Records what the route asked of the vector store."""

    def __init__(self) -> None:
        self.indexed: list[dict[str, Any]] = []
        self.deleted: list[dict[str, Any]] = []

    async def index_pages(self, **kwargs: Any) -> int:
        self.indexed.append(kwargs)
        return len(kwargs["pages"])

    async def delete_file(self, **kwargs: Any) -> None:
        self.deleted.append(kwargs)


@dataclass
class Vault:
    client: AsyncClient
    fake: FakeStore
    storage: Path
    _holder: dict[str, object] = field(default_factory=dict)

    def as_user(self, user: object) -> AsyncClient:
        self._holder["user"] = user
        return self.client


def _build_app(
    session_factory, holder: dict[str, object], *, max_body_bytes: int | None = None
) -> FastAPI:
    from app.api import files as files_module
    from app.core.deps import current_user, get_session

    app = FastAPI()
    app.add_middleware(
        BodySizeLimitMiddleware,
        max_bytes=max_body_bytes or (settings.max_upload_mb + 1) * 1024 * 1024,
    )
    app.include_router(files_module.router, prefix="/api/v1")

    async def request_session():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = request_session
    app.dependency_overrides[current_user] = lambda: holder["user"]
    return app


@pytest_asyncio.fixture
async def vault(session_factory, tmp_path, monkeypatch):
    from app.api import files as files_module
    from app.rag import store

    fake = FakeStore()
    monkeypatch.setattr(store, "index_pages", fake.index_pages)
    monkeypatch.setattr(store, "delete_file", fake.delete_file)
    storage = tmp_path / "uploads"
    monkeypatch.setattr(settings, "storage_dir", str(storage))
    files_module.upload_rate_limiter.reset()

    holder: dict[str, object] = {}
    app = _build_app(session_factory, holder)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield Vault(client=client, fake=fake, storage=storage, _holder=holder)
    app.dependency_overrides.clear()
    files_module.upload_rate_limiter.reset()


@pytest_asyncio.fixture
async def alice(db):
    org = await make_org(db, "acme")
    user = await make_user(db, org, "alice@acme.test")
    await db.commit()
    return user


async def upload(client: AsyncClient, name: str, payload: bytes, **form: str):
    return await client.post(
        "/api/v1/files",
        files={"file": (name, payload, "application/octet-stream")},
        data=form,
    )


def _stored_files(storage: Path) -> list[Path]:
    return [p for p in storage.rglob("*") if p.is_file()]


async def _asset_count(session_factory) -> int:
    async with session_factory() as session:
        return (await session.execute(select(func.count(FileAsset.id)))).scalar_one()


# ──────────────────────────── filename handling ────────────────────────────


def test_sanitize_filename_strips_paths_and_control_chars():
    assert sanitize_filename("../../../etc/passwd.pdf") == "passwd.pdf"
    assert sanitize_filename("..\\..\\share\\evil.pdf") == "evil.pdf"
    assert sanitize_filename("report\r\nX-Injected: yes.pdf") == "reportX-Injected: yes.pdf"
    assert sanitize_filename("quarterly‮report.pdf") == "quarterlyreport.pdf"
    assert sanitize_filename("...hidden.pdf") == "hidden.pdf"
    assert sanitize_filename(None) == ""
    assert sanitize_filename("..") == ""
    long = sanitize_filename("a" * 400 + ".pdf")
    assert len(long) <= 255 and long.endswith(".pdf")


async def test_traversal_name_lands_as_uuid_on_disk(vault, alice, session_factory):
    response = await upload(
        vault.as_user(alice), "../../../etc/cron.d/evil.pdf", minimal_pdf()
    )
    assert response.status_code == 201
    body = response.json()
    assert body["filename"] == "evil.pdf"
    assert body["indexed"] is True

    stored = _stored_files(vault.storage)
    assert len(stored) == 1
    # The disk name is {owner_id}/{asset_id}{ext} — no user-supplied part.
    assert stored[0].name == f"{body['id']}.pdf"
    assert stored[0].parent.name == str(alice.id)
    assert stored[0].resolve().is_relative_to(vault.storage.resolve())
    # The citation name in the index is the sanitised one.
    assert vault.fake.indexed[0]["filename"] == "evil.pdf"


# ─────────────────────────── content vs extension ──────────────────────────


async def test_zip_wearing_a_pdf_name_is_refused(vault, alice, session_factory):
    response = await upload(vault.as_user(alice), "slides.pdf", plain_zip())
    assert response.status_code == 415
    assert _stored_files(vault.storage) == []
    assert await _asset_count(session_factory) == 0


async def test_pdf_wearing_a_docx_name_is_refused(vault, alice):
    response = await upload(vault.as_user(alice), "notes.docx", minimal_pdf())
    assert response.status_code == 415


async def test_binary_wearing_a_text_name_is_refused(vault, alice):
    response = await upload(vault.as_user(alice), "notes.txt", b"\x89PNG\x00\x1a\nbinary")
    assert response.status_code == 415


async def test_zip_wearing_a_docx_name_fails_indexing_with_a_clear_error(vault, alice):
    """A real zip passes the magic check; the parser inside the sandbox
    refuses it, and the refusal lands on the asset, not on the API."""
    response = await upload(vault.as_user(alice), "totally-a-doc.docx", plain_zip())
    assert response.status_code == 201
    body = response.json()
    assert body["indexed"] is False
    assert body["index_error"]
    assert vault.fake.indexed == []


async def test_client_content_type_is_never_stored(vault, alice, session_factory):
    response = await vault.as_user(alice).post(
        "/api/v1/files",
        files={"file": ("a.pdf", minimal_pdf(), 'text/html"><script>')},
    )
    assert response.status_code == 201
    async with session_factory() as session:
        asset = (await session.execute(select(FileAsset))).scalar_one()
    assert asset.content_type == "application/pdf"


async def test_empty_file_is_refused(vault, alice):
    response = await upload(vault.as_user(alice), "empty.pdf", b"")
    assert response.status_code == 422


async def test_unsupported_extension_is_refused(vault, alice):
    response = await upload(vault.as_user(alice), "shell.sh", b"#!/bin/sh\n")
    assert response.status_code == 415


# ────────────────────────────── size limits ────────────────────────────────


async def test_oversize_stream_is_cut_off_with_no_leftovers(
    vault, alice, session_factory, monkeypatch
):
    monkeypatch.setattr(settings, "max_upload_mb", 1)
    payload = b"%PDF-" + b"A" * (1024 * 1024 + 64)
    response = await upload(vault.as_user(alice), "big.pdf", payload)
    assert response.status_code == 413
    assert _stored_files(vault.storage) == []  # no final file, no .part
    assert await _asset_count(session_factory) == 0


async def test_declared_oversize_body_is_refused_at_the_transport(vault, alice):
    """Larger than the middleware ceiling: refused before anything runs."""
    payload = b"A" * ((settings.max_upload_mb + 2) * 1024 * 1024)
    response = await upload(vault.as_user(alice), "huge.pdf", payload)
    assert response.status_code == 413
    assert _stored_files(vault.storage) == []


async def test_chunked_body_without_content_length_is_cut_off(session_factory, alice):
    """A lying (chunked) client is stopped mid-multipart, as a clean 413."""
    holder: dict[str, object] = {"user": alice}
    # A tiny transport cap keeps the test fast.
    app = _build_app(session_factory, holder, max_body_bytes=1024)

    boundary = "e5boundary"
    part = (
        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; '
        f'filename="big.pdf"\r\nContent-Type: application/pdf\r\n\r\n'
    ).encode() + b"%PDF-" + b"A" * 4096 + f"\r\n--{boundary}--\r\n".encode()

    async def dribble():
        for i in range(0, len(part), 256):
            yield part[i : i + 256]

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/api/v1/files",
            content=dribble(),
            headers={"content-type": f"multipart/form-data; boundary={boundary}"},
        )
    assert response.status_code == 413


async def test_body_limit_passes_ordinary_requests(vault, alice):
    response = await upload(vault.as_user(alice), "fine.pdf", minimal_pdf())
    assert response.status_code == 201


# ─────────────────────────── parser containment ────────────────────────────


async def test_malformed_pdf_records_error_and_api_survives(vault, alice):
    response = await upload(
        vault.as_user(alice), "broken.pdf", b"%PDF-1.7\nnot really a pdf at all"
    )
    assert response.status_code == 201
    body = response.json()
    assert body["indexed"] is False
    assert body["index_error"]

    # The poisoned file took nothing down: the next upload indexes fine.
    again = await upload(vault.as_user(alice), "fine.docx", docx_bytes())
    assert again.status_code == 201
    assert again.json()["indexed"] is True


async def test_page_budget_is_enforced(vault, alice, monkeypatch):
    monkeypatch.setattr(settings, "max_pages_per_file", 3)
    response = await upload(vault.as_user(alice), "deck.pptx", pptx_bytes(slides=5))
    assert response.status_code == 201
    body = response.json()
    assert body["indexed"] is False
    assert "limit is 3" in body["index_error"]
    assert vault.fake.indexed == []


async def test_parse_timeout_is_recorded_not_raised(vault, alice, monkeypatch):
    monkeypatch.setattr(settings, "parse_timeout_seconds", 0)
    response = await upload(vault.as_user(alice), "slow.docx", docx_bytes())
    assert response.status_code == 201
    body = response.json()
    assert body["indexed"] is False
    assert "did not finish" in body["index_error"]


async def test_sandbox_parses_real_files(tmp_path):
    for name, payload, needle in (
        ("a.pdf", minimal_pdf("Verification matrix."), "Verification matrix."),
        ("a.docx", docx_bytes(("First paragraph.",)), "First paragraph."),
        ("a.pptx", pptx_bytes(slides=2), "Slide 1 content"),
    ):
        path = tmp_path / name
        path.write_bytes(payload)
        pages = await parse_in_sandbox(path, path.suffix)
        assert pages and needle in pages[0].text


async def test_sandbox_timeout_raises_parse_timeout(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "parse_timeout_seconds", 0)
    path = tmp_path / "a.docx"
    path.write_bytes(docx_bytes())
    with pytest.raises(ParseTimeout):
        await parse_in_sandbox(path, ".docx")


async def test_sandbox_rejects_spoofed_container_with_safe_message(tmp_path):
    path = tmp_path / "a.docx"
    path.write_bytes(plain_zip())
    with pytest.raises(ParseRejected) as excinfo:
        await parse_in_sandbox(path, ".docx")
    assert ".docx" in str(excinfo.value)


# ────────────────────────────── rate limiting ──────────────────────────────


async def test_upload_rate_limit_answers_429(vault, alice, monkeypatch):
    from app.api import files as files_module

    monkeypatch.setattr(files_module.upload_rate_limiter, "max_attempts", 2)
    ok1 = await upload(vault.as_user(alice), "a.docx", docx_bytes())
    ok2 = await upload(vault.as_user(alice), "b.docx", docx_bytes())
    blocked = await upload(vault.as_user(alice), "c.docx", docx_bytes())
    assert (ok1.status_code, ok2.status_code, blocked.status_code) == (201, 201, 429)


# ─────────────────────────────── no orphans ────────────────────────────────


async def test_commit_failure_leaves_no_bytes_and_no_points(
    engine, session_factory, tmp_path, monkeypatch, alice
):
    """If the row cannot be committed, the disk bytes and the just-written
    Qdrant points are removed before the error surfaces."""
    from app.rag import store
    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

    class ExplodingCommit(AsyncSession):
        async def commit(self):
            raise RuntimeError("boom")

    exploding_factory = async_sessionmaker(
        engine, expire_on_commit=False, class_=ExplodingCommit
    )

    fake = FakeStore()
    monkeypatch.setattr(store, "index_pages", fake.index_pages)
    monkeypatch.setattr(store, "delete_file", fake.delete_file)
    storage = tmp_path / "uploads"
    monkeypatch.setattr(settings, "storage_dir", str(storage))

    holder: dict[str, object] = {"user": alice}
    app = _build_app(exploding_factory, holder)
    transport = ASGITransport(app=app, raise_app_exceptions=False)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/api/v1/files",
            files={"file": ("a.docx", docx_bytes(), "application/octet-stream")},
        )

    assert response.status_code == 500
    assert _stored_files(storage) == []
    assert len(fake.indexed) == 1  # points were written…
    assert len(fake.deleted) == 1  # …and cleaned up again
    assert fake.deleted[0]["file_id"] == fake.indexed[0]["file_id"]
    assert await _asset_count(session_factory) == 0


# ──────────────────── listing follows the retrieval scope ──────────────────


@pytest_asyncio.fixture
async def world(db):
    """acme: alice owns p_file (in project) and u_file (unfiled); bob has no
    grants. globex: eve owns g_file."""
    from .factories import make_file_asset

    acme = await make_org(db, "acme")
    globex = await make_org(db, "globex")
    alice = await make_user(db, acme, "alice@acme.test")
    bob = await make_user(db, acme, "bob@acme.test")
    eve = await make_user(db, globex, "eve@globex.test")
    project = await make_project(db, acme, alice, name="Tender")
    await make_grant(
        db, org_id=acme.id, subject_type="project", subject_id=project.id,
        user=alice, role="owner",
    )
    p_file = await make_file_asset(db, alice, "tender.pdf", project=project)
    u_file = await make_file_asset(db, alice, "memo.pdf")
    g_file = await make_file_asset(db, eve, "rival.pdf")
    await db.commit()
    return {
        "acme": acme, "alice": alice, "bob": bob, "eve": eve,
        "project": project, "p_file": p_file, "u_file": u_file, "g_file": g_file,
    }


async def _listed_ids(client: AsyncClient) -> set[str]:
    response = await client.get("/api/v1/files")
    assert response.status_code == 200
    return {row["id"] for row in response.json()}


async def test_listing_shows_own_files(vault, world):
    assert await _listed_ids(vault.as_user(world["alice"])) == {
        str(world["p_file"].id), str(world["u_file"].id),
    }


async def test_listing_shows_nothing_without_a_grant(vault, world):
    assert await _listed_ids(vault.as_user(world["bob"])) == set()


async def test_listing_admits_project_shared_files_only(vault, world, db):
    await make_grant(
        db, org_id=world["acme"].id, subject_type="project",
        subject_id=world["project"].id, user=world["bob"], role="viewer",
    )
    await db.commit()
    # The project file appears; alice's unfiled memo never does.
    assert await _listed_ids(vault.as_user(world["bob"])) == {str(world["p_file"].id)}


async def test_listing_never_crosses_the_org_boundary(vault, world):
    assert await _listed_ids(vault.as_user(world["eve"])) == {str(world["g_file"].id)}


async def test_shared_visibility_does_not_confer_delete(vault, world, db, session_factory):
    await make_grant(
        db, org_id=world["acme"].id, subject_type="project",
        subject_id=world["project"].id, user=world["bob"], role="viewer",
    )
    await db.commit()
    response = await vault.as_user(world["bob"]).delete(
        f"/api/v1/files/{world['p_file'].id}"
    )
    assert response.status_code == 404
    assert vault.fake.deleted == []
    async with session_factory() as session:
        assert (
            await session.get(FileAsset, world["p_file"].id)
        ) is not None


async def test_role_change_requires_ownership(vault, world, db):
    await make_grant(
        db, org_id=world["acme"].id, subject_type="project",
        subject_id=world["project"].id, user=world["bob"], role="viewer",
    )
    await db.commit()
    response = await vault.as_user(world["bob"]).patch(
        f"/api/v1/files/{world['p_file'].id}/role", params={"role": "target"}
    )
    assert response.status_code == 404


# A session id belonging to someone else must still 404 (E4's check, kept).
async def test_foreign_session_id_is_refused(vault, world, db):
    from app.db.models import ChatSession

    foreign = ChatSession(
        owner_id=world["eve"].id, org_id=world["eve"].org_id, title="theirs"
    )
    db.add(foreign)
    await db.commit()
    response = await vault.as_user(world["alice"]).post(
        "/api/v1/files",
        files={"file": ("a.docx", docx_bytes(), "application/octet-stream")},
        data={"session_id": str(foreign.id)},
    )
    assert response.status_code == 404
